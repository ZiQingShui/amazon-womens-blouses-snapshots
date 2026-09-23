#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""从 ExpandKeywords 导出表里挑「核心词」——用于每日关键词排名采集。

⚠ 2026-09-23 起，**权威清单改为人工指定的 `config/keywords-core.json`（14 个词）**，
   本脚本降级为「备选/扩充」用途：只有当用户要重新按规则扩词时才跑，产物是 work/keywords-core.json，
   不会自动被采集脚本读取（`fetch_keyword_search.py` 已默认指向 config/）。
   部署前请确认 `config/keywords-core.json` 才是要被采集的那份。

为什么不能只按搜索量取 Top N：表里搜索量最大的词是 womens dresses / pants for women /
sweaters for women 这类**别的品类**，抓到它们的前 5 名对本项目没有意义。所以先按女式衬衫品类
（shirt / blouse / top / tunic / tee / button-down …）过滤，再按月搜索量排序。

用法：
  python tools/pick_keywords.py --xlsx "work/ExpandKeywords-US-xxx.xlsx" --limit 150 --out work/keywords-core.json
"""
from __future__ import annotations

import argparse
import json
import pathlib
import re
import sys

ROOT = pathlib.Path(__file__).resolve().parent.parent

# 女式衬衫品类相关词
KEEP = re.compile(r"\b(shirts?|blouses?|tops?|tunics?|tees?|t-shirts?|blouson|poplin|chiffon|"
                  r"button[- ]?downs?|button[- ]?ups?|shackets?)\b", re.I)
# 明确别的品类 —— 命中就丢（注意 \b 边界：shorts 不会误伤 shirt）
DROP = re.compile(r"\b(dresses?|pants?|skirts?|jeans|denim|sweaters?|cardigans?|jackets?|coats?|"
                  r"shoes?|boots?|sandals?|flats?|heels?|bags?|purses?|hats?|scarfs?|scarves|"
                  r"leggings?|shorts|rompers?|jumpsuits?|bodysuits?|bras?|underwears?|socks?|"
                  r"suits?|blazers?|hoodies?|sweatshirts?|tanks?|camis?|swimsuits?|bikinis?|"
                  r"pajamas?|sleepwear|robes?|costumes?|socks|corsets?|bralettes?|bustiers?)\b", re.I)


def num(x):
    try:
        s = str(x).replace(",", "").strip()
        return float(s) if s not in ("", "None", "nan") else None
    except Exception:
        return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--limit", type=int, default=150)
    ap.add_argument("--out", default=str(ROOT / "work/keywords-core.json"))
    ap.add_argument("--min-volume", type=int, default=0, help="低于这个月搜索量的不要")
    ap.add_argument("--min-words", type=int, default=2,
                    help="最少单词数，默认 2 —— 单词词（top / tops / blouse）搜索量再大也不要，"
                         "Amazon 给的结果会跑偏、对选品没有指导意义")
    args = ap.parse_args()

    import openpyxl
    wb = openpyxl.load_workbook(args.xlsx, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    head = next(it)
    rows = [r for r in it if r and r[0]]

    # ⚠ 曾有一版用「列8 是否含『相关性』」来跳过 ExpandKeywords 的示例行 —— **那是错的**：
    # 实测列8 只有前 3 行有值，而那 3 行是**真实关键词**（`white button down shirt women` 12.5 万、
    # `button down shirts for women` 13.4 万、`womens tops` 115 万），只是被人工加了相关性标注。
    # 按那条规则跳会把 12~13 万搜索量的词整条丢掉（2026-09-23 由用户给的清单暴露出来）。
    # 现在一律不丢，是否入选完全交给下面的 KEEP / DROP / 词数 / 搜索量规则。

    cand, dropped_kind, no_volume, too_short = [], 0, 0, 0
    for r in rows:
        kw = str(r[0]).strip()
        if not kw:
            continue
        if DROP.search(kw) or not KEEP.search(kw):
            dropped_kind += 1
            continue
        if len(re.findall(r"[A-Za-z0-9']+", kw)) < args.min_words:
            too_short += 1
            continue
        vol = num(r[3])
        if vol is None:
            no_volume += 1
            continue
        if vol < args.min_volume:
            continue
        cand.append({
            "keyword": kw,
            "zh": str(r[1] or "").strip(),
            "monthlyVolume": int(vol),
            "abaRank": int(num(r[2])) if num(r[2]) is not None else None,
            "supplyDemand": num(r[4]),
        })

    cand.sort(key=lambda x: -x["monthlyVolume"])
    picked = cand[: args.limit]
    out = pathlib.Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps({
        "schemaVersion": 1,
        "source": pathlib.Path(args.xlsx).name,
        "rule": "含 shirt/blouse/top/tunic/tee/button-down 且不含 dress/pants/sweater 等别品类；按月搜索量降序",
        "totalWords": len(rows), "candidates": len(cand), "picked": len(picked),
        "keywords": picked,
    }, ensure_ascii=False, indent=1), encoding="utf-8")

    print("总词 %d → 品类过滤后 %d（丢别品类 %d / 单词词 %d / 无搜索量 %d）→ 取前 %d"
          % (len(rows), len(cand), dropped_kind, too_short, no_volume, len(picked)))
    print("月搜索量区间：%d ~ %d" % (picked[-1]["monthlyVolume"], picked[0]["monthlyVolume"]) if picked else "（空）")
    print()
    for i, k in enumerate(picked[:25], 1):
        print("  %3d. %-44s %8d  ABA %s" % (i, k["keyword"][:44], k["monthlyVolume"], k["abaRank"]))
    if len(picked) > 25:
        print("  ... 还有 %d 个" % (len(picked) - 25))
    print("\n已写出", out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
